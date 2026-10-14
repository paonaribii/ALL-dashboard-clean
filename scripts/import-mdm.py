"""Import an "All Status MDM" workbook into contracts.db and re-match every contract.

Usage:  python scripts/import-mdm.py "C:/path/All Status MDM.xlsx"

Each sheet has IMEI (or Apple serial) / STATUS / Brand columns. Rows with a blank status are skipped.
A "No MDM" status in the workbook is a confirmed result (the device was looked up), so it counts as
matched; contracts whose device is in no MDM source at all keep a NULL status.

Every imported workbook is kept in the mdm_imports table (one batch per run), so earlier imports are not
lost. Priority when a device appears in several sources: newest import > older imports > mdm_new_raw >
mdm_status (oldest). Any sheet whose name contains "Apple" is treated as the Apple (serial) sheet.
Apple contracts look up the Apple sheet, other brands the IMEI sheets; both the contract's IMEI and
serial fields are tried. Afterwards run `npm run build:data`.
"""
import io
import sqlite3
import sys
from pathlib import Path

import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
DB = Path(__file__).resolve().parent.parent / 'data' / 'contracts.db'
STATUS_MAP = {'SUCCESS': 'MDM Success', 'UNSUCCESS': 'MDM Unsuccess', 'NO MDM': 'No MDM'}


def norm_key(v):
    if v is None:
        return None
    if isinstance(v, float):
        return None if v != v else str(int(v))
    if isinstance(v, int):
        return str(v)
    return str(v).strip().upper() or None


rows = []
wb = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
for sheet in wb.sheetnames:
    for r in wb[sheet].iter_rows(min_row=2, values_only=True):
        if not r or len(r) < 2:
            continue
        key, status = norm_key(r[0]), STATUS_MAP.get(str(r[1] or '').strip().upper())
        if key and status:
            rows.append((key, 'Apple' if 'apple' in sheet.lower() else sheet, status))
wb.close()
print('rows imported:', len(rows))

conn = sqlite3.connect(DB)
cur = conn.cursor()
cur.execute('CREATE TABLE IF NOT EXISTS mdm_imports (device_key TEXT, brand_sheet TEXT, status TEXT, batch INTEGER, source TEXT)')
if not cur.execute('SELECT 1 FROM mdm_imports LIMIT 1').fetchone() and cur.execute(
        "SELECT 1 FROM sqlite_master WHERE name='mdm_latest_raw'").fetchone():
    # The first version of this script kept only the last workbook (mdm_latest_raw); keep it as batch 1.
    cur.execute("INSERT INTO mdm_imports SELECT device_key, brand_sheet, status, 1, 'All Status MDM (2).xlsx' FROM mdm_latest_raw")
    cur.execute('DROP TABLE mdm_latest_raw')
batch = cur.execute('SELECT COALESCE(MAX(batch), 0) + 1 FROM mdm_imports').fetchone()[0]
cur.executemany('INSERT INTO mdm_imports VALUES (?,?,?,?,?)', [(*r, batch, Path(sys.argv[1]).name) for r in rows])
print('import batch:', batch)

before = dict(cur.execute("SELECT COALESCE(status,'(not found)'), COUNT(*) FROM contract_mdm GROUP BY 1"))

RANK = "CASE status WHEN 'MDM Success' THEN 1 WHEN 'MDM Unsuccess' THEN 2 ELSE 3 END"
cur.execute('DROP TABLE IF EXISTS mdm_keyed')
cur.execute(f'''CREATE TABLE mdm_keyed AS
  SELECT device_key, is_apple, status FROM (
    SELECT device_key, is_apple, status,
      ROW_NUMBER() OVER (PARTITION BY device_key, is_apple ORDER BY prio, {RANK}) rn
    FROM (SELECT device_key, brand_sheet='Apple' is_apple, status, -batch prio FROM mdm_imports
          UNION ALL SELECT device_key, brand_sheet='Apple' is_apple, status, 1 prio FROM mdm_new_raw
          UNION ALL SELECT device_key, brand_sheet='Apple' is_apple, status, 2 prio FROM mdm_status)
  ) WHERE rn = 1''')
cur.execute('CREATE INDEX idx_mdm_keyed ON mdm_keyed(is_apple, device_key)')

cur.execute('DROP TABLE IF EXISTS contract_mdm')
cur.execute('''CREATE TABLE contract_mdm AS
SELECT c.id,
  (SELECT m.status FROM mdm_keyed m
   WHERE m.is_apple = (UPPER(TRIM(c.brand_name)) LIKE 'APPLE%')
     AND m.device_key IN (UPPER(TRIM(c.imei)), UPPER(TRIM(c.serial_number)))
   ORDER BY CASE m.status WHEN 'MDM Success' THEN 1 WHEN 'MDM Unsuccess' THEN 2 ELSE 3 END LIMIT 1) AS status
FROM contracts c''')
cur.execute('CREATE UNIQUE INDEX idx_contract_mdm ON contract_mdm(id)')
conn.commit()

after = dict(cur.execute("SELECT COALESCE(status,'(not found)'), COUNT(*) FROM contract_mdm GROUP BY 1"))
total = sum(after.values())
for k in sorted(set(before) | set(after)):
    print(f'  {k:14} {before.get(k, 0):>8,} -> {after.get(k, 0):>8,}')
found = total - after.get('(not found)', 0)
print(f'matched: {found:,} / {total:,} = {found / total * 100:.1f}%')
conn.close()
